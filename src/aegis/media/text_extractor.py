"""
Text extractor — converts document/tabular/code/data files to plain text.
Output is a clean string ready to inject into LLM context.
"""
import io
import json
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

MAX_CHARS = 12_000  # ~3000 tokens — safe context budget per file


def extract(data: bytes, filename: str) -> str:
    """
    Extract readable text from any non-image file.
    Returns clean text string, truncated to MAX_CHARS.
    """
    ext = Path(filename).suffix.lower()

    try:
        if ext == ".pdf":
            return _pdf(data)
        elif ext in (".docx", ".doc"):
            return _docx(data)
        elif ext in (".xlsx", ".xls"):
            return _excel(data)
        elif ext == ".csv":
            return _csv(data)
        elif ext == ".tsv":
            return _csv(data, sep="\t")
        elif ext == ".json":
            return _json(data)
        elif ext == ".xml":
            return _xml(data)
        elif ext in (".txt", ".md", ".rtf"):
            return _plain(data)
        else:
            # Code files and anything else — read as plain text
            return _plain(data)
    except Exception as e:
        logger.warning(f"TextExtractor: failed for '{filename}': {e}")
        return f"[Could not extract text from '{filename}': {e}]"



def _pdf(data: bytes) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(f"[Page {i+1}]\n{text}")
            return _trim("\n\n".join(pages))
    except ImportError:
        # Fallback to pypdf
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        return _trim(text)


def _docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return _trim(text)


def _excel(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"[Sheet: {sheet}]\n" + "\n".join(rows))
    return _trim("\n\n".join(parts))


def _csv(data: bytes, sep: str = ",") -> str:
    import csv
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text), delimiter=sep)
    rows = [" | ".join(row) for row in reader if any(c.strip() for c in row)]
    return _trim("\n".join(rows))


def _json(data: bytes) -> str:
    try:
        obj = json.loads(data.decode("utf-8", errors="replace"))
        return _trim(json.dumps(obj, indent=2, ensure_ascii=False))
    except Exception:
        return _plain(data)


def _xml(data: bytes) -> str:
    import xml.etree.ElementTree as ET
    try:
        root = ET.fromstring(data)
        lines = []
        for elem in root.iter():
            if elem.text and elem.text.strip():
                lines.append(f"{elem.tag}: {elem.text.strip()}")
        return _trim("\n".join(lines))
    except Exception:
        return _plain(data)


def _plain(data: bytes) -> str:
    return _trim(data.decode("utf-8", errors="replace"))


def _trim(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + f"\n\n[... truncated — showing first {MAX_CHARS} characters]"
    return text
